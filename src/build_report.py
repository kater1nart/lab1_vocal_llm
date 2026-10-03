"""Сборка отчёта report/*.md -> report/Отчет_ЛР1.docx (заголовки, абзацы, списки, таблицы, код, картинки)."""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
REP = ROOT / "report"


def add_runs(par, text):
    # **жирный**, `код`, [текст](ссылка) -> "текст (ссылка)"
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = "Courier New"
            r.font.size = Pt(10)
        elif tok.startswith("*") and len(tok) > 2:
            par.add_run(tok[1:-1]).italic = True
        else:
            par.add_run(tok)


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-+:?", c) for c in r)]
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Table Grid"
    for i, r in enumerate(cells):
        for j, c in enumerate(r[:len(cells[0])]):
            p = t.cell(i, j).paragraphs[0]
            add_runs(p, c)
            for run in p.runs:
                run.font.size = Pt(9)
                run.bold = run.bold or i == 0
    doc.add_paragraph()


def convert(md, doc):
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            j = i + 1
            while not lines[j].startswith("```"):
                j += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(lines[i + 1:j]))
            r.font.name = "Courier New"
            r.font.size = Pt(8)
            i = j + 1
            continue
        if ln.startswith("|"):
            j = i
            while j < len(lines) and lines[j].startswith("|"):
                j += 1
            add_table(doc, lines[i:j])
            i = j
            continue
        m = re.match(r"!\[(.*)\]\((.+)\)", ln)
        if m:
            doc.add_picture(str((REP / m.group(2)).resolve()), width=Cm(16))
            cap = doc.add_paragraph(m.group(1))
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif ln.startswith("#"):
            level = len(ln) - len(ln.lstrip("#"))
            doc.add_heading(ln.lstrip("#").strip(), level=min(level, 3))
        elif re.match(r"^\s*[-*] ", ln):
            add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", ln))
        elif re.match(r"^\s*\d+\. ", ln):
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\. ", "", ln))
        elif ln.strip():
            add_runs(doc.add_paragraph(), ln)
        i += 1


def main():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    for f in sorted(REP.glob("[0-9][0-9]_*.md")):
        convert(f.read_text(), doc)
    out = REP / "Отчет_ЛР1.docx"
    doc.save(out)
    print("saved", out)


if __name__ == "__main__":
    main()
